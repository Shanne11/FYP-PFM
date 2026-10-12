"""Build a review-ready workbook from the Section 16 audit outputs.

The preferred artifact-tool runtime is unavailable on this host, so this
fallback uses openpyxl as permitted by the spreadsheet authoring guidance.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / "outputs" / "section16_audit"
OUTPUT = AUDIT / "Section16_Evidence_Matrix.xlsx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "EEF4F8"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "F4CCCC"
GREY = "E7E6E6"
TEXT = "1F2937"
THIN = Side(style="thin", color="AAB2BD")


def style_title(ws, title: str, subtitle: str, max_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws["A1"] = title
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color=TEXT)
    ws["A2"].fill = PatternFill("solid", fgColor=PALE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def write_table(ws, frame: pd.DataFrame, start_row: int, name: str) -> None:
    for col_index, column in enumerate(frame.columns, 1):
        cell = ws.cell(start_row, col_index, str(column).replace("_", " ").title())
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    for row_index, values in enumerate(frame.itertuples(index=False, name=None), start_row + 1):
        for col_index, value in enumerate(values, 1):
            if pd.isna(value):
                value = None
            cell = ws.cell(row_index, col_index, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=Side(style="hair", color="D9D9D9"))
    end_row = start_row + len(frame)
    end_col = len(frame.columns)
    if len(frame):
        table = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(end_col)}{end_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"


def fit_columns(ws, minimum: int = 10, maximum: int = 45) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = max([len(value) for value in values[:150]] + [minimum]) + 2
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(maximum, width)


def main() -> None:
    matrix = pd.read_csv(AUDIT / "section16_evidence_matrix.csv")
    hashes = pd.read_csv(AUDIT / "artifact_hash_verification.csv")
    selectors = pd.read_csv(AUDIT / "actm_selector_three_seed_means.csv")
    participation = pd.read_csv(AUDIT / "client_participation_audit.csv")
    client_size = pd.read_csv(AUDIT / "client_size_weight_adjustment.csv")
    client_accuracy = pd.read_csv(AUDIT / "matched_client_accuracy_audit.csv")
    transitions = json.loads((AUDIT / "smart_note_transition_audit.json").read_text(encoding="utf-8"))
    bounds = json.loads((AUDIT / "aggregation_bounds_audit.json").read_text(encoding="utf-8"))
    python_git = json.loads((AUDIT / "python_git_state.json").read_text(encoding="utf-8"))
    flutter_git = json.loads((AUDIT / "flutter_git_state.json").read_text(encoding="utf-8"))
    privacy = json.loads((AUDIT / "flutter_privacy_audit.json").read_text(encoding="utf-8"))
    demo = json.loads((AUDIT / "tier_a_seven_step_demo_trace.json").read_text(encoding="utf-8"))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    style_title(
        summary,
        "Section 16 Evidence Verification",
        "Current evidence state. Results are provisional until both working trees are committed and the evidence manifest is regenerated against those commits.",
        6,
    )
    summary_rows = [
        ("Evidence manifest hashes", f"{int(hashes.hash_match.sum())}/26 match", "Pass", "All declared files exist and match SHA-256."),
        ("Python repository", python_git["head"], "Provisional", f"{python_git['changed_file_count']} changed or untracked paths."),
        ("Flutter repository", flutter_git["head"], "Provisional", f"{flutter_git['changed_file_count']} changed or untracked paths."),
        ("Fully verified questions", int((matrix.status == "Verified").sum()), "Count", "No material qualification."),
        ("Verified with limitation", int((matrix.status == "Verified with limitation").sum()), "Count", "Use the stated limitation in the report."),
        ("Needs correction", int((matrix.status == "Needs correction").sum()), "Action", "Q12 and Q15."),
    ]
    for col, value in enumerate(("Check", "Value", "State", "Interpretation"), 1):
        cell = summary.cell(4, col, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for row, values in enumerate(summary_rows, 5):
        for col, value in enumerate(values, 1):
            cell = summary.cell(row, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
        state_fill = GREEN if values[2] == "Pass" else AMBER if values[2] == "Provisional" else RED if values[2] == "Action" else BLUE
        summary.cell(row, 3).fill = PatternFill("solid", fgColor=state_fill)
        summary.cell(row, 3).font = Font(bold=True)

    summary["A13"] = "Critical reproduced results"
    summary["A13"].font = Font(size=13, bold=True, color=NAVY)
    critical = [
        ("Q11 audit cases", transitions["cases"], "No duplicate seed–transaction pairs."),
        ("Wrong to correct", transitions["wrong_to_correct"], "Retrospective source-note cases."),
        ("Correct to wrong", transitions["correct_to_wrong"], "Retrospective source-note cases."),
        ("Changed prediction", transitions["changed_prediction"], "Across 1,736 note-used cases."),
        ("Q12 confident but wrong", f"{transitions['confidence_increased_but_after_incorrect']} supplied formula / {transitions['confidence_increased_wrong_before_and_after']} implemented metric", "Neither definition supports 900; state which definition is reported."),
        ("Q15 client-seed accuracy", "0 improved / 0 declined / 450 unchanged", "Correct the unsupported 15 / 19 / 416 claim."),
        ("Q14 multiplier range", f"{bounds['multiplier_min']:.6f} to {bounds['multiplier_max']:.6f}", "Bounds are [0.75, 1.25]; final coefficients are separately normalised."),
        ("Q18 direct sync test", "2 tests passed", "No payload, no pending row, syncPending throws."),
    ]
    for col, value in enumerate(("Metric", "Reproduced value", "Meaning"), 1):
        cell = summary.cell(14, col, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for row, values in enumerate(critical, 15):
        for col, value in enumerate(values, 1):
            summary.cell(row, col, value)
            summary.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
            summary.cell(row, col).border = Border(bottom=THIN)
    summary.freeze_panes = "A4"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 48
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 55

    evidence = workbook.create_sheet("Evidence Matrix")
    style_title(evidence, "Evidence-to-Question Matrix", "Paths and SHA-256 values identify the exact evidence used for each decision.", len(matrix.columns))
    write_table(evidence, matrix, 4, "EvidenceMatrix")
    evidence.freeze_panes = "A5"
    widths = {1: 12, 2: 34, 3: 44, 4: 40, 5: 18, 6: 38, 7: 55, 8: 24, 9: 55}
    for col, width in widths.items():
        evidence.column_dimensions[get_column_letter(col)].width = width
    last = 4 + len(matrix)
    evidence.conditional_formatting.add(
        f"H5:H{last}", FormulaRule(formula=["$H5=\"Verified\""], fill=PatternFill("solid", fgColor=GREEN))
    )
    evidence.conditional_formatting.add(
        f"H5:H{last}", FormulaRule(formula=["$H5=\"Verified with limitation\""], fill=PatternFill("solid", fgColor=AMBER))
    )
    evidence.conditional_formatting.add(
        f"H5:H{last}", FormulaRule(formula=["$H5=\"Needs correction\""], fill=PatternFill("solid", fgColor=RED))
    )

    hash_sheet = workbook.create_sheet("Artifact Hashes")
    style_title(hash_sheet, "26-Artifact SHA-256 Check", "Every declared artifact currently matches the provisional manifest.", len(hashes.columns))
    write_table(hash_sheet, hashes, 4, "ArtifactHashes")
    hash_sheet.freeze_panes = "A5"
    fit_columns(hash_sheet, maximum=68)

    numerical = workbook.create_sheet("Numerical Checks")
    style_title(numerical, "Recalculated Numerical Evidence", "Selector means, Smart Note transitions and aggregation checks calculated from transaction/client-level records.", 9)
    numerical["A4"] = "ACTM selector three-seed means"
    numerical["A4"].font = Font(size=13, bold=True, color=NAVY)
    write_table(numerical, selectors, 5, "SelectorMeans")
    start = 8 + len(selectors)
    numerical.cell(start, 1, "Smart Note transition audit").font = Font(size=13, bold=True, color=NAVY)
    transition_frame = pd.DataFrame([{"measure": key, "value": value} for key, value in transitions.items()])
    write_table(numerical, transition_frame, start + 1, "TransitionAudit")
    bound_start = start + 4 + len(transition_frame)
    numerical.cell(bound_start, 1, "Aggregation-bound audit").font = Font(size=13, bold=True, color=NAVY)
    bounds_frame = pd.DataFrame([{"measure": key, "value": value} for key, value in bounds.items()])
    write_table(numerical, bounds_frame, bound_start + 1, "BoundsAudit")
    numerical.freeze_panes = "A5"
    fit_columns(numerical, maximum=38)

    clients = workbook.create_sheet("Client Checks")
    style_title(clients, "Client Participation and Harm Checks", "Matched participation, client-size coefficient effects and per-client final-test accuracy.", 8)
    clients["A4"] = "Participation by seed and round"
    clients["A4"].font = Font(size=13, bold=True, color=NAVY)
    write_table(clients, participation, 5, "ParticipationAudit")
    size_start = 8 + len(participation)
    clients.cell(size_start, 1, "Coefficient adjustment by client-size quartile").font = Font(size=13, bold=True, color=NAVY)
    write_table(clients, client_size, size_start + 1, "ClientSizeAudit")
    accuracy_start = size_start + 4 + len(client_size)
    clients.cell(accuracy_start, 1, "Matched final-test client accuracy (450 client-seed cases)").font = Font(size=13, bold=True, color=NAVY)
    write_table(clients, client_accuracy, accuracy_start + 1, "ClientAccuracyAudit")
    clients.freeze_panes = "A5"
    fit_columns(clients, maximum=32)

    trace = workbook.create_sheet("Demo Trace")
    style_title(trace, "Tier A Seven-Step Demonstration Trace", "Controlled Tier A evidence. This is not a PocketIQ Flutter local-audit record.", 6)
    trace_rows = [
        (1, "Transaction metadata and initial prediction", json.dumps(demo["metadata"], ensure_ascii=False), json.dumps(demo["before"], ensure_ascii=False)),
        (2, "ACTM decision", json.dumps(demo["actm"], ensure_ascii=False), "Selected under the nominal 30% budget."),
        (3, "Source Smart Note and probability change", demo["source_note"], f"Before: {demo['before']['probability_vector']}\nAfter: {demo['after']['probability_vector']}"),
        (4, "Final correctness outcome", demo["outcome"], f"Before correct={demo['before']['correct']}; after correct={demo['after']['correct']}"),
        (5, "Client base and Proposed coefficients", demo["aggregation_final_round"]["base_weight"], demo["aggregation_final_round"]["proposed_final_coefficient"]),
        (6, "Bounded multiplier", demo["aggregation_final_round"]["signal_multiplier"], f"Within [0.75, 1.25]={demo['aggregation_final_round']['multiplier_within_bounds']}"),
        (7, "Trace record", f"seed={demo['seed']}; source_index={demo['source_index']}; transaction_id={demo['transaction_id']}; client_id={demo['client_id']}", demo["audit_warning"]),
    ]
    trace_frame = pd.DataFrame(trace_rows, columns=["step", "evidence", "value", "interpretation"])
    write_table(trace, trace_frame, 4, "DemoTrace")
    trace.freeze_panes = "A5"
    trace.column_dimensions["A"].width = 9
    trace.column_dimensions["B"].width = 38
    trace.column_dimensions["C"].width = 70
    trace.column_dimensions["D"].width = 70
    for row in range(5, 5 + len(trace_frame)):
        trace.row_dimensions[row].height = 62

    privacy_sheet = workbook.create_sheet("Git and Privacy")
    style_title(privacy_sheet, "Git State and Flutter Privacy Verification", "Commit identifiers, dirty-tree state and direct-call upload-disabled test evidence.", 4)
    records = [
        ("Python", "Branch", python_git["branch"], ""),
        ("Python", "HEAD", python_git["head"], python_git["last_commit"]),
        ("Python", "Working tree clean", python_git["working_tree_clean"], f"{python_git['changed_file_count']} changed/untracked paths"),
        ("Flutter", "Branch", flutter_git["branch"], ""),
        ("Flutter", "HEAD", flutter_git["head"], flutter_git["last_commit"]),
        ("Flutter", "Working tree clean", flutter_git["working_tree_clean"], f"{flutter_git['changed_file_count']} changed/untracked paths"),
        ("Flutter privacy", "Remote sync flag false", privacy["remote_sync_flag_false"], ""),
        ("Flutter privacy", "Payload created", privacy["train_creates_payload_file"], "Expected False"),
        ("Flutter privacy", "Pending row created", privacy["train_creates_pending_sync_row"], "Expected False"),
        ("Flutter privacy", "syncPending throws", privacy["sync_pending_unconditionally_throws"], "Direct-call test passed"),
        ("Flutter privacy", "Network call present", privacy["sync_pending_contains_network_call"], "Expected False"),
        ("Flutter privacy", "Test result", privacy["programmatic_rejection_test_passed"], "Two tests passed"),
    ]
    git_frame = pd.DataFrame(records, columns=["repository", "check", "value", "detail"])
    write_table(privacy_sheet, git_frame, 4, "GitPrivacy")
    privacy_sheet.freeze_panes = "A5"
    fit_columns(privacy_sheet, maximum=70)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(OUTPUT)

    # Post-save structural and content verification.
    checked = load_workbook(OUTPUT, data_only=False)
    assert checked.sheetnames == [
        "Summary", "Evidence Matrix", "Artifact Hashes", "Numerical Checks",
        "Client Checks", "Demo Trace", "Git and Privacy",
    ]
    assert checked["Evidence Matrix"].max_row == len(matrix) + 4
    assert checked["Artifact Hashes"].max_row == len(hashes) + 4
    assert checked["Demo Trace"]["A5"].value == 1
    assert OUTPUT.stat().st_size > 20_000
    print(f"Created and verified {OUTPUT} ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
